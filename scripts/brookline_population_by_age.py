"""Build Brookline population-by-age history + projections, and resident births.

Inputs (all under data/demographics/, see output/brookline_population_by_age_notes.md):
  census2000_sf1/brookline_2000_seg02.txt   SF1 2000 file 02 row, LOGRECNO 0078543 (P12 at data offset 121)
  census2010_sf1/brookline_2010_seg04.txt   SF1 2010 file 04 row, LOGRECNO 0110482 (P12 at data offset 144)
  census2020_dhc/brookline_2020_seg06.txt   DHC 2020 segment 06 row, LOGRECNO 0000315 (P12 at data offset 144)
  acs5yr/brookline_acs5y{2019,2024}_b01001.txt  ACS 5-year table-based summary file rows
  umdi_projections/*.xlsx                    UMDI V2024 and UMDI-DOT V2018 projections
  births/*                                   DLS databank (1980-2017) and DPH annual birth reports

Outputs:
  output/brookline_population_by_age.csv
  output/brookline_resident_births.csv
"""
import csv
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEM = os.path.join(ROOT, "data", "demographics")
OUT = os.path.join(ROOT, "output")

# P12 / B01001 cell order (identical across 2000 SF1, 2010 SF1, 2020 DHC and ACS B01001;
# confirmed from Census API group metadata p12_*.json and the ACS table shells).
P12_AGES = ["0-4", "5-9", "10-14", "15-17", "18-19", "20", "21", "22-24", "25-29", "30-34",
            "35-39", "40-44", "45-49", "50-54", "55-59", "60-61", "62-64", "65-66", "67-69",
            "70-74", "75-79", "80-84", "85+"]

FIVE_YR = ["0-4", "5-9", "10-14", "15-19", "20-24", "25-29", "30-34", "35-39", "40-44", "45-49",
           "50-54", "55-59", "60-64", "65-69", "70-74", "75-79", "80-84", "85+"]

# five-year group -> P12 detail cells
FIVE_FROM_P12 = {
    "15-19": ["15-17", "18-19"], "20-24": ["20", "21", "22-24"],
    "60-64": ["60-61", "62-64"], "65-69": ["65-66", "67-69"],
}

# Summary groups requested for the project. Those needing 15-17 / 18-19 detail
# exist only for Census/ACS; UMDI (five-year groups only) gets the 5-14 / 15-19 variants.
SUMMARY_DETAIL = {
    "under 5": ["0-4"],
    "5-17": ["5-9", "10-14", "15-17"],
    "18-24": ["18-19", "20", "21", "22-24"],
}
SUMMARY_FIVE = {
    "5-14": ["5-9", "10-14"],
    "5-19": ["5-9", "10-14", "15-19"],
    "25-44": ["25-29", "30-34", "35-39", "40-44"],
    "45-64": ["45-49", "50-54", "55-59", "60-64"],
    "65-74": ["65-69", "70-74"],
    "75-84": ["75-79", "80-84"],
    "85+": ["85+"],
    "65+": ["65-69", "70-74", "75-79", "80-84", "85+"],
    "75+": ["75-79", "80-84", "85+"],
}

KNOWN_TOTALS = {2000: 57107, 2010: 58732, 2020: 63191}


def p12_from_row(path, sep, offset):
    """Return 49 P12 cells from a one-line summary-file extract, after validating structure."""
    fields = open(path).read().strip().split(sep)
    data = fields[5:]  # FILEID, STUSAB, CHARITER, CIFSN, LOGRECNO precede the data cells
    v = [int(x) for x in data[offset:offset + 49]]
    assert v[1] + v[25] == v[0], "male + female != total"
    assert sum(v[2:25]) == v[1], "male age cells do not sum to male total"
    assert sum(v[26:49]) == v[25], "female age cells do not sum to female total"
    return v


def detail_from_p12(v):
    """Both-sexes population for each P12 detail age cell."""
    return {a: v[2 + i] + v[26 + i] for i, a in enumerate(P12_AGES)}


def acs_b01001(path):
    lines = open(path).read().strip().split("\n")
    d = dict(zip(lines[0].split("|"), lines[1].split("|")))
    assert d["GEO_ID"] == "0600000US2502109175"
    v = [int(d["B01001_E%03d" % i]) for i in range(1, 50)]
    m = [int(d["B01001_M%03d" % i]) for i in range(1, 50)]
    assert v[1] + v[25] == v[0] and sum(v[2:25]) == v[1] and sum(v[26:49]) == v[25]
    return v, m


def rows_from_detail(year, source, det, total, is_proj):
    five = {}
    for g in FIVE_YR:
        five[g] = sum(det[c] for c in FIVE_FROM_P12.get(g, [g]))
    assert sum(five.values()) == total
    out = [(year, source, "total", total, is_proj, "total")]
    out += [(year, source, g, five[g], is_proj, "five_year") for g in FIVE_YR]
    for g, cells in SUMMARY_DETAIL.items():
        out.append((year, source, g, sum(det[c] for c in cells), is_proj, "summary"))
    for g, cells in SUMMARY_FIVE.items():
        out.append((year, source, g, sum(five[c] for c in cells), is_proj, "summary"))
    return out


def rows_from_five(year, source, five, total, is_proj):
    out = [(year, source, "total", total, is_proj, "total")]
    out += [(year, source, g, five[g], is_proj, "five_year") for g in FIVE_YR]
    out.append((year, source, "under 5", five["0-4"], is_proj, "summary"))
    for g, cells in SUMMARY_FIVE.items():
        out.append((year, source, g, sum(five[c] for c in cells), is_proj, "summary"))
    return out


def umdi_v2024():
    f = os.path.join(DEM, "umdi_projections",
                     "UMDI_V2024_Long-Term_Population_Projections_MCD,_County,_RPA,_State_Age_Sex_detail_2010-2050.xlsx")
    ws = openpyxl.load_workbook(f, read_only=True)["CityTown Age 2010-2050"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    br = [r for r in rows[1:] if r[0] == "Brookline"]
    assert [r[6] for r in br] == FIVE_YR
    res = {}
    for j, y in enumerate(hdr):
        if isinstance(y, int):
            res[y] = {r[6]: int(r[j]) for r in br}
    return res


def umdi_v2018():
    f = os.path.join(DEM, "umdi_projections", "AgeSexDetails_UMDI-DOT_V2018.xlsx")
    ws = openpyxl.load_workbook(f, read_only=True)["CityTown_Age"]
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            hdr = r
        elif r[0] == "BROOKLINE":
            rows.append(r)
        elif rows:
            break
    assert [r[6] for r in rows] == FIVE_YR
    res = {}
    for j, h in enumerate(hdr):
        if isinstance(h, str) and (h.startswith("Projection") or h.startswith("Census")):
            res[int(h.split()[-1])] = {r[6]: int(r[j]) for r in rows}
    return res


def main():
    out = []
    census = {
        2000: ("census2000_sf1/brookline_2000_seg02.txt", ",", 121, "Census 2000 SF1 P12"),
        2010: ("census2010_sf1/brookline_2010_seg04.txt", ",", 144, "Census 2010 SF1 P12"),
        2020: ("census2020_dhc/brookline_2020_seg06.txt", "|", 144, "Census 2020 DHC P12"),
    }
    census_five = {}
    for y, (p, sep, off, src) in census.items():
        v = p12_from_row(os.path.join(DEM, p), sep, off)
        assert v[0] == KNOWN_TOTALS[y], (y, v[0])
        det = detail_from_p12(v)
        rs = rows_from_detail(y, src, det, v[0], False)
        census_five[y] = {r[2]: r[3] for r in rs if r[5] == "five_year"}
        out += rs

    for y, period in [(2019, "2015-2019"), (2024, "2020-2024")]:
        v, _ = acs_b01001(os.path.join(DEM, "acs5yr", "brookline_acs5y%d_b01001.txt" % y))
        out += rows_from_detail(y, "ACS 5-year %s B01001 (estimate)" % period, detail_from_p12(v), v[0], False)

    u24 = umdi_v2024()
    # UMDI 2010/2020 columns are the Census counts (Brookline had no launch adjustment); verify, don't duplicate.
    for y in (2010, 2020):
        assert u24[y] == census_five[y], "UMDI V2024 %d column differs from Census" % y
    for y in sorted(u24):
        if y > 2020:
            out += rows_from_five(y, "UMDI V2024 projection", u24[y], sum(u24[y].values()), True)

    u18 = umdi_v2018()
    assert u18[2010] == census_five[2010], "UMDI V2018 2010 column differs from Census"
    for y in sorted(u18):
        if y > 2010:
            out += rows_from_five(y, "UMDI-DOT V2018 projection", u18[y], sum(u18[y].values()), True)

    os.makedirs(OUT, exist_ok=True)
    with open(os.path.join(OUT, "brookline_population_by_age.csv"), "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["year", "source", "age_group", "population", "is_projection", "grouping"])
        w.writerows(out)

    # Resident births. DPH annual reports are authoritative where available; DLS databank otherwise.
    dls = dict(zip(range(1980, 2018), [433, 486, 493, 533, 545, 553, 572, 536, 618, 621, 649, 605, 584, 625,
                                       585, 600, 808, 599, 655, 589, 614, 599, 649, 691, 697, 613, 691, 668,
                                       666, 676, 693, 688, 686, 718, 653, 666, 605, 594]))
    dph = {2015: 609, 2016: 605, 2017: 594, 2018: 571, 2019: 581, 2020: 487, 2021: 491, 2022: 482,
           2023: 422}   # 2023: Table 17 of the 2023 report (data/MA-DPH-2023-Birth-Report.docx); NESDEC had estimated 518
    with open(os.path.join(OUT, "brookline_resident_births.csv"), "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["year", "resident_births", "source", "note"])
        for y in range(1980, 2024):
            if y in dph:
                note = ""
                if y in dls and dls[y] != dph[y]:
                    note = "DLS databank shows %d; DPH report value used" % dls[y]
                w.writerow([y, dph[y], "MA DPH Massachusetts Births %d, municipal resident births table" % y, note])
            else:
                note = "possible data-entry outlier; not cross-checked" if y == 1996 else ""
                w.writerow([y, dls[y], "MA DOR DLS Municipal Databank, Number of Resident Births", note])
    print("rows:", len(out))


if __name__ == "__main__":
    main()

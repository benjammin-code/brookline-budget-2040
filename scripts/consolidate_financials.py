"""
Consolidate Brookline's historical financials (FY2005-2025) into one clean table:
  - Levy-limit mechanics (prior levy limit, 2.5% increase, new growth, override,
    debt exclusions, excess capacity, total levy) from Ben's levy-history model.
  - Schedule A General Fund revenue by function, all years, Brookline only.
  - Schedule A General Fund expenditure by function, all years, Brookline only.

  - DOR Recap-based revenue by source (Tax Levy / State Aid / Local Receipts /
    All Other, without Enterprise & CPA funds) for cross-validation.
  - DOR-certified residential vs. total new growth split.
  - Local Receipts actual-vs-estimate totals (FY2017-2025 only).

Also writes a second, separate file with category-level Local Receipts
actual-vs-estimate detail (34 receipt types x FY2017-2025) since that's too
granular to fold into the one-row-per-year master table.

Sources: see ../data/SOURCES.md
Outputs:
  ../output/brookline_historical_financials_fy05_25.csv
  ../output/brookline_local_receipts_detail_fy17_25.csv
"""
import csv
import glob
import re

import openpyxl
import xlrd

DATA_DIR = "../data"
SCHED_A_DIR = f"{DATA_DIR}/revenue & expenditures all MA"
LEVY_FILE = f"{DATA_DIR}/Brookline property tax revenue over 20 years.xlsx"
REV_BY_SOURCE_FILE = f"{SCHED_A_DIR}/RevBySource 2005-2025.xlsx"
NEW_GROWTH_FILE = f"{DATA_DIR}/2026 MA state new growth for all towns 20 years.xlsx"
OUT_FILE = "../output/brookline_historical_financials_fy05_25.csv"
LOCAL_RECEIPTS_DETAIL_OUT = "../output/brookline_local_receipts_detail_fy17_25.csv"

REV_COLS = [
    "taxes", "service_charges", "licenses_permits", "federal_revenue",
    "state_revenue", "revenue_other_govts", "special_assessments",
    "fines_forfeitures", "miscellaneous", "other_financing_sources",
    "transfers", "total_revenues",
]
EXP_COLS = [
    "general_government", "public_safety", "education", "public_works",
    "human_services", "culture_recreation", "fixed_costs",
    "intergov_assessments", "other_expenditures", "debt_service",
    "total_expenditures",
]


def load_schedule_a(kind, cols):
    """kind: 'Revenues' or 'Expenditures'. Returns {year: {col: value}}."""
    out = {}
    for f in glob.glob(f"{SCHED_A_DIR}/GenFund{kind}*.xlsx"):
        m = re.search(r"(\d{4})\.xlsx$", f)
        year = int(m.group(1))
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb["Sheet1"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[1] == "Brookline":
                values = row[3:]  # skip DOR code, municipality, fiscal year
                out[year] = dict(zip(cols, values))
                break
        else:
            raise ValueError(f"Brookline row not found in {f}")
    return out


def load_levy_history():
    """Returns {year: {...levy mechanics...}} for actual years FY2005-2025."""
    wb = openpyxl.load_workbook(LEVY_FILE, data_only=True)
    ws = wb["sources"]
    out = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            continue
        year = int(row[0])
        if year > 2025:
            continue  # forecast years, not actuals -- out of scope for this table
        out[year] = {
            "prior_levy_limit": row[1],
            "increase_2_5pct": row[2],
            "new_growth": row[3],
            "override": row[4],
            "debt_exclusions": row[5],
            "excess_capacity": row[6],
            "total_levy": row[7],
        }
    return out


def load_rev_by_source():
    """DOR Recap-based Tax Levy / State Aid / Local Receipts / All Other,
    without Enterprise & CPA funds. Returns {year: {...}}."""
    wb = openpyxl.load_workbook(REV_BY_SOURCE_FILE, data_only=True)
    ws = wb.worksheets[0]
    out = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[1] == "Brookline":
            year = int(row[2])
            out[year] = {
                "recap_tax_levy": row[3],
                "recap_state_aid": row[4],
                "recap_local_receipts": row[5],
                "recap_all_other": row[6],
                "recap_total_no_enterprise": row[8],
            }
    return out


def load_new_growth_split():
    """Residential vs total new growth (DOR-certified), all years. Returns {year: {...}}."""
    wb = openpyxl.load_workbook(NEW_GROWTH_FILE, data_only=True)
    ws = wb["data"]
    out = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[1] == "Brookline":
            year = int(row[2])
            out[year] = {
                "res_new_growth_value": row[3],
                "res_new_growth_applied": row[4],
                "total_new_growth_value": row[5],
                "total_new_growth_applied": row[6],
                "res_new_growth_pct_of_total": row[7],
            }
    return out


def load_local_receipts_detail():
    """Category-level Local Receipts estimate vs. actual, FY2017-2025 only
    (not available further back). Returns (totals: {year: {est, act}},
    detail_rows: [(year, receipt_type_id, description, estimate, actual)])."""
    totals = {}
    detail_rows = []
    for f in glob.glob(f"{SCHED_A_DIR}/LocalReciptesActVsEst *.xls"):
        year = int(re.search(r"(\d{4})", f).group(1))
        wb = xlrd.open_workbook(f)
        ws = wb.sheet_by_index(0)
        est_total = act_total = 0
        for r in range(1, ws.nrows):
            row = ws.row_values(r)
            if len(row) > 1 and row[1] == "Brookline":
                _, _, _, receipt_id, desc, est, act = row
                est_total += est
                act_total += act
                detail_rows.append((year, int(receipt_id), desc, est, act))
        totals[year] = {"local_receipts_est_total": est_total,
                         "local_receipts_act_total": act_total}
    return totals, detail_rows


def main():
    revenues = load_schedule_a("Revenues", REV_COLS)
    expenditures = load_schedule_a("Expenditures", EXP_COLS)
    levy = load_levy_history()
    rev_by_source = load_rev_by_source()
    new_growth = load_new_growth_split()
    local_receipts_totals, local_receipts_detail = load_local_receipts_detail()

    years = sorted(set(revenues) & set(expenditures) & set(levy)
                   & set(rev_by_source) & set(new_growth))
    missing_rev = sorted(set(levy) - set(revenues))
    missing_exp = sorted(set(levy) - set(expenditures))
    if missing_rev:
        print(f"WARNING: levy-history years with no Schedule A revenue match: {missing_rev}")
    if missing_exp:
        print(f"WARNING: levy-history years with no Schedule A expenditure match: {missing_exp}")

    levy_cols = ["prior_levy_limit", "increase_2_5pct", "new_growth", "override",
                 "debt_exclusions", "excess_capacity", "total_levy"]
    recap_cols = ["recap_tax_levy", "recap_state_aid", "recap_local_receipts",
                  "recap_all_other", "recap_total_no_enterprise"]
    ng_cols = ["res_new_growth_value", "res_new_growth_applied",
               "total_new_growth_value", "total_new_growth_applied",
               "res_new_growth_pct_of_total"]
    lr_cols = ["local_receipts_est_total", "local_receipts_act_total"]
    header = (["fiscal_year"] + levy_cols
              + [f"rev_{c}" for c in REV_COLS]
              + [f"exp_{c}" for c in EXP_COLS]
              + recap_cols + ng_cols + lr_cols
              + ["general_fund_surplus", "taxes_vs_total_levy_gap"])

    rows = []
    for y in years:
        lv = levy[y]
        rv = revenues[y]
        ex = expenditures[y]
        rbs = rev_by_source[y]
        ng = new_growth[y]
        # Local Receipts actual-vs-estimate detail only available FY2017-2025;
        # earlier years get blank cells rather than shrinking the whole table.
        lr = local_receipts_totals.get(y, {c: "" for c in lr_cols})
        surplus = rv["total_revenues"] - ex["total_expenditures"]
        tax_gap = rv["taxes"] - lv["total_levy"]
        row = ([y] + [lv[c] for c in levy_cols]
               + [rv[c] for c in REV_COLS]
               + [ex[c] for c in EXP_COLS]
               + [rbs[c] for c in recap_cols]
               + [ng[c] for c in ng_cols]
               + [lr[c] for c in lr_cols]
               + [surplus, tax_gap])
        rows.append(row)

    with open(OUT_FILE, "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(rows)

    print(f"Wrote {len(rows)} years ({years[0]}-{years[-1]}) to {OUT_FILE}")

    print("\ntaxes_vs_total_levy_gap (Schedule A 'Taxes' minus levy-model 'total_levy'):")
    for row in rows:
        y, gap = row[0], row[-1]
        print(f"  {y}: {gap:>14,.0f}")

    with open(LOCAL_RECEIPTS_DETAIL_OUT, "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["fiscal_year", "receipt_type_id", "receipt_description",
                          "estimate", "actual"])
        writer.writerows(sorted(local_receipts_detail))
    lr_years = sorted(local_receipts_totals)
    print(f"\nWrote {len(local_receipts_detail)} rows "
          f"({lr_years[0]}-{lr_years[-1]}) to {LOCAL_RECEIPTS_DETAIL_OUT}")


if __name__ == "__main__":
    main()
